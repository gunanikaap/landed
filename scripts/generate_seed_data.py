"""
generate_seed_data.py
---------------------------------
Generates a realistic but fully SYNTHETIC job-search pipeline.
No real people or private data — safe to commit and deploy publicly.

Outputs (relative to repo root):
  sample_data/companies.csv
  sample_data/applications.csv
  sample_data/contacts.csv
  sample_data/events.csv
  job_applications_sample.xlsx   (all four as separate sheets, for viewing)

Swap this out for your real tracker later — keep the same column names
and the rest of the app (DuckDB load + agent) keeps working unchanged.
"""

import os
import random
from datetime import date, timedelta

import pandas as pd

SEED = 42
random.seed(SEED)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "sample_data")
os.makedirs(DATA_DIR, exist_ok=True)

TODAY = date(2026, 6, 3)

# ---------------------------------------------------------------- companies
COMPANIES = [
    ("Apple", "Technology", "Enterprise", "Cork"),
    ("Google", "Technology", "Enterprise", "Dublin"),
    ("Mastercard", "FinTech", "Enterprise", "Dublin"),
    ("Stripe", "FinTech", "Large", "Dublin"),
    ("Intercom", "Technology", "Mid", "Dublin"),
    ("Workday", "Technology", "Large", "Dublin"),
    ("HubSpot", "Technology", "Large", "Dublin"),
    ("Fidelity Investments", "Financial Services", "Enterprise", "Dublin"),
    ("BNY", "Financial Services", "Enterprise", "Cork"),
    ("Optum", "Healthcare", "Enterprise", "Dublin"),
    ("Deloitte", "Consulting", "Enterprise", "Dublin"),
    ("Accenture", "Consulting", "Enterprise", "Dublin"),
    ("Version 1", "Consulting", "Large", "Dublin"),
    ("Red Hat", "Technology", "Large", "Waterford"),
    ("Microsoft", "Technology", "Enterprise", "Dublin"),
    ("Amazon (AWS)", "Technology", "Enterprise", "Dublin"),
    ("Salesforce", "Technology", "Enterprise", "Dublin"),
    ("LinkedIn", "Technology", "Large", "Dublin"),
    ("Datadog", "Technology", "Large", "Dublin"),
    ("Liberty IT", "Insurance", "Mid", "Dublin"),
    ("Aon", "Insurance", "Enterprise", "Dublin"),
    ("Guidewire", "InsurTech", "Mid", "Dublin"),
    ("Fenergo", "FinTech", "Mid", "Dublin"),
    ("Phorest", "Technology", "SME", "Dublin"),
    ("Squarespace", "Technology", "Large", "Dublin"),
    ("SAP", "Technology", "Enterprise", "Dublin"),
    ("Indeed", "Technology", "Large", "Dublin"),
    ("Zalando", "Technology", "Large", "Dublin"),
]

companies = []
for i, (name, industry, size, loc) in enumerate(COMPANIES, start=1):
    companies.append({
        "company_id": i,
        "company_name": name,
        "industry": industry,
        "size_band": size,
        "hq_location": loc,
        "has_irish_office": "Y",
    })
companies_df = pd.DataFrame(companies)
company_ids = companies_df["company_id"].tolist()

# ---------------------------------------------------------------- lookups
ROLE_BY_TRACK = {
    "Data Engineer": ["Data Engineer", "AWS Data Engineer", "Senior Data Engineer",
                      "Data Engineer (Analytics)", "AI / Data Engineer"],
    "Product Owner": ["Product Owner", "Technical Product Owner",
                      "Product Owner (Insurance)", "Associate Product Owner"],
    "Scrum Master": ["Scrum Master", "Agile Delivery Lead", "Senior Scrum Master"],
    "Data Analyst": ["Data Analyst", "Analytics Engineer", "BI Analyst"],
    "Business Analyst": ["Business Analyst", "Data Business Analyst"],
}
TRACK_WEIGHTS = [("Data Engineer", 0.40), ("Product Owner", 0.25),
                 ("Scrum Master", 0.12), ("Data Analyst", 0.15),
                 ("Business Analyst", 0.08)]
SENIORITY = ["Associate", "Mid", "Senior"]
CHANNELS = [("Referral", 0.30), ("Cold Apply", 0.45),
            ("Recruiter", 0.18), ("Inbound", 0.07)]
SOURCE_BY_CHANNEL = {
    "Referral": ["LinkedIn", "Internal Referral"],
    "Cold Apply": ["Company Site", "Indeed", "LinkedIn"],
    "Recruiter": ["Recruiter Agency", "LinkedIn"],
    "Inbound": ["LinkedIn", "Recruiter Agency"],
}
WORK_MODE = [("Hybrid", 0.55), ("Remote", 0.30), ("Onsite", 0.15)]

# A realistic funnel. Each "stage reached" implies all earlier events happened.
# (status, current_stage, weight)
OUTCOMES = [
    ("Ghosted", "No Response", 0.26),
    ("Rejected", "Application Review", 0.20),
    ("Screening", "Recruiter Screen", 0.12),
    ("Rejected", "Recruiter Screen", 0.08),
    ("Interview", "Hiring Manager Call", 0.08),
    ("Interview", "Technical Interview", 0.07),
    ("Rejected", "Technical Interview", 0.06),
    ("Take-home", "Take-home Submitted", 0.04),
    ("Interview", "Final / Onsite", 0.05),
    ("Offer", "Offer", 0.05),
    ("Withdrawn", "Withdrawn", 0.03),
    ("Applied", "Application Review", 0.00),  # buffer
]

STAGE_ORDER = [
    "Applied", "Recruiter Screen", "Hiring Manager Call",
    "Technical Interview", "Take-home Submitted", "Final / Onsite", "Offer",
]


def weighted_choice(pairs):
    r, cum = random.random(), 0.0
    for value, w in pairs:
        cum += w
        if r <= cum:
            return value
    return pairs[-1][0]


def salary_band(track, seniority):
    base = {"Data Engineer": 70, "Product Owner": 68, "Scrum Master": 65,
            "Data Analyst": 55, "Business Analyst": 58}[track]
    bump = {"Associate": -8, "Mid": 0, "Senior": 14}[seniority]
    lo = (base + bump) * 1000
    hi = lo + random.choice([8000, 10000, 12000, 15000])
    target = lo + random.choice([3000, 5000, 7000])
    return lo, hi, target


# ---------------------------------------------------------------- applications + events
applications, events = [], []
event_id = 1
N_APPS = 56

for app_id in range(1, N_APPS + 1):
    cid = random.choice(company_ids)
    track = weighted_choice(TRACK_WEIGHTS)
    role = random.choice(ROLE_BY_TRACK[track])
    seniority = random.choices(SENIORITY, weights=[3, 5, 2])[0]
    channel = weighted_choice(CHANNELS)
    source = random.choice(SOURCE_BY_CHANNEL[channel])
    mode = weighted_choice(WORK_MODE)
    loc = companies_df.loc[companies_df.company_id == cid, "hq_location"].iloc[0]
    location = "Remote (Ireland)" if mode == "Remote" else loc
    lo, hi, target = salary_band(track, seniority)

    days_ago = random.randint(2, 110)
    applied = TODAY - timedelta(days=days_ago)

    status, stage, _ = random.choices(OUTCOMES, weights=[o[2] for o in OUTCOMES])[0]
    if status == "Applied":
        status, stage = "Ghosted", "No Response"

    # Build the event trail up to the stage reached
    reached = stage if stage in STAGE_ORDER else "Applied"
    trail = ["Applied"]
    for s in STAGE_ORDER[1:]:
        trail.append(s)
        if s == reached:
            break
    if stage in ("No Response", "Withdrawn", "Application Review"):
        trail = ["Applied"]

    cursor = applied
    last_activity = applied
    for ev in trail:
        cursor = cursor + timedelta(days=random.randint(2, 9))
        if cursor > TODAY:
            cursor = TODAY
        events.append({
            "event_id": event_id,
            "application_id": app_id,
            "event_type": ev,
            "event_date": cursor.isoformat(),
            "channel": channel,
            "notes": "",
        })
        event_id += 1
        last_activity = max(last_activity, cursor)

    # Terminal event for closed outcomes
    if status in ("Rejected", "Offer", "Withdrawn"):
        cursor = min(cursor + timedelta(days=random.randint(2, 12)), TODAY)
        events.append({
            "event_id": event_id, "application_id": app_id,
            "event_type": status, "event_date": cursor.isoformat(),
            "channel": channel, "notes": "",
        })
        event_id += 1
        last_activity = max(last_activity, cursor)

    applications.append({
        "application_id": app_id,
        "company_id": cid,
        "role_title": role,
        "track": track,
        "seniority": seniority,
        "channel": channel,
        "source": source,
        "work_mode": mode,
        "location": location,
        "date_applied": applied.isoformat(),
        "status": status,
        "current_stage": stage,
        "salary_min": lo,
        "salary_max": hi,
        "salary_target": target,
        "last_activity_date": last_activity.isoformat(),
    })

applications_df = pd.DataFrame(applications)
events_df = pd.DataFrame(events)

# ---------------------------------------------------------------- contacts (anonymised)
FIRST = ["Aoife", "Liam", "Niamh", "Sean", "Saoirse", "Cian", "Emer", "Darragh",
         "Ciara", "Oisin", "Roisin", "Eoin", "Maeve", "Conor", "Sinead", "Fionn",
         "Orla", "Padraig", "Aisling", "Cormac", "Grainne", "Tadhg"]
LAST = ["Byrne", "Murphy", "Kelly", "O'Brien", "Walsh", "Ryan", "O'Connor",
        "Doyle", "McCarthy", "Gallagher", "Healy", "Lynch", "Brennan", "Fitzgerald",
        "Nolan", "Quinn", "Daly", "Moran", "Kavanagh", "Dunne", "Carroll", "Hayes"]
REL = [("Referral", 0.4), ("Recruiter", 0.3), ("Hiring Manager", 0.15), ("Network", 0.15)]
TITLE_BY_REL = {
    "Referral": ["Senior Data Engineer", "Product Owner", "Engineering Manager", "Scrum Master"],
    "Recruiter": ["Talent Acquisition Partner", "Technical Recruiter", "Sourcing Specialist"],
    "Hiring Manager": ["Engineering Manager", "Head of Data", "Delivery Lead"],
    "Network": ["Data Analyst", "Software Engineer", "Maynooth Alumni"],
}

contacts = []
random.shuffle(FIRST)
n_contacts = 22
for j in range(1, n_contacts + 1):
    cid = random.choice(company_ids)
    rel = weighted_choice(REL)
    name = f"{FIRST[j % len(FIRST)]} {random.choice(LAST)}"
    connected = TODAY - timedelta(days=random.randint(5, 120))
    responded = "Y" if random.random() < (0.7 if rel != "Network" else 0.4) else "N"
    contacts.append({
        "contact_id": j,
        "company_id": cid,
        "full_name": name,
        "title": random.choice(TITLE_BY_REL[rel]),
        "relationship": rel,
        "connected_on": connected.isoformat(),
        "responded": responded,
    })
contacts_df = pd.DataFrame(contacts)

# ---------------------------------------------------------------- write CSVs
companies_df.to_csv(os.path.join(DATA_DIR, "companies.csv"), index=False)
applications_df.to_csv(os.path.join(DATA_DIR, "applications.csv"), index=False)
contacts_df.to_csv(os.path.join(DATA_DIR, "contacts.csv"), index=False)
events_df.to_csv(os.path.join(DATA_DIR, "events.csv"), index=False)

# ---------------------------------------------------------------- write viewable workbook
xlsx_path = os.path.join(ROOT, "job_applications_sample.xlsx")
with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xl:
    applications_df.to_excel(xl, sheet_name="applications", index=False)
    companies_df.to_excel(xl, sheet_name="companies", index=False)
    contacts_df.to_excel(xl, sheet_name="contacts", index=False)
    events_df.to_excel(xl, sheet_name="events", index=False)

# light formatting: bold header + sane column widths
from openpyxl import load_workbook
from openpyxl.styles import Font

wb = load_workbook(xlsx_path)
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Calibri")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 34)
    ws.freeze_panes = "A2"
wb.save(xlsx_path)

print("companies   :", len(companies_df))
print("applications:", len(applications_df))
print("contacts    :", len(contacts_df))
print("events      :", len(events_df))
print("status mix  :")
print(applications_df["status"].value_counts().to_string())
print("\nwrote:", xlsx_path)
print("wrote CSVs to:", DATA_DIR)
